import io
import hashlib
import json
import shutil
import struct
import subprocess
import time
import unittest
import wave
import zipfile
import zlib
from fastapi.testclient import TestClient
import os
import tempfile
_TEST_DATA = tempfile.TemporaryDirectory()
os.environ['STUDIO_DATA_DIR'] = _TEST_DATA.name
from services.api import main


def png():
    def chunk(kind,content):
        return struct.pack('>I',len(content))+kind+content+struct.pack('>I',zlib.crc32(kind+content)&0xffffffff)
    return b'\x89PNG\r\n\x1a\n'+chunk(b'IHDR',struct.pack('>IIBBBBB',32,32,8,2,0,0,0))+chunk(b'IDAT',zlib.compress((b'\0'+bytes([100,150,90])*32)*32))+chunk(b'IEND',b'')

def wav():
    out=io.BytesIO()
    with wave.open(out,'wb') as w:
        w.setnchannels(1);w.setsampwidth(2);w.setframerate(22050);w.writeframes(b'\0\0'*22050)
    return out.getvalue()

class ProductionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.context=TestClient(main.app);cls.client=cls.context.__enter__()
        cls.client.post('/api/v1/auth/login',json={'password':'test-only-password-123'})
    @classmethod
    def tearDownClass(cls):cls.context.__exit__(None,None,None)
    def project(self):return self.client.post('/api/v1/projects',json={'name':'Production CPU'}).json()
    def media(self,p,blob=None,mime='image/png'):
        r=self.client.post('/api/v1/assets/upload',params={'project_id':p['id']},files={'file':('test.png' if mime=='image/png' else 'test.wav',blob or png(),mime)})
        self.assertEqual(r.status_code,201);return r.json()
    def state(self,p):return self.client.get('/api/v1/projects/'+p['id']+'/production').json()
    def wait_render(self,p,rid):
        deadline=time.monotonic()+30
        while time.monotonic()<deadline:
            row=next(r for r in self.state(p)['renders'] if r['id']==rid)
            if row['status'] in ('completed','failed','canceled'):return row
            time.sleep(.1)
        self.fail('Render did not finish within 30 seconds')
    def test_script_split_version_and_preservation(self):
        p=self.project();url='/api/v1/projects/'+p['id']
        r=self.client.put(url+'/script',json={'text':'Première scène.\n\nالمشهد الثاني','version':0})
        self.assertEqual(r.status_code,200)
        self.assertEqual(self.client.put(url+'/script',json={'text':'stale','version':0}).status_code,409)
        shots=self.client.post(url+'/script/split').json()['shots']
        self.assertEqual(len(shots),2)
        self.assertEqual(self.client.post(url+'/script/split').status_code,409)
        self.assertEqual(self.state(p)['script']['text'],'Première scène.\n\nالمشهد الثاني')
        a=self.media(p)
        body={**shots[0],'title':'Ouverture','asset_id':a['id'],'approved':True,'duration':1}
        r=self.client.put(url+'/shots/'+shots[0]['id'],json=body)
        self.assertEqual(r.status_code,200)
        self.assertEqual(r.json()['version'],2)
        self.assertEqual(self.client.put(url+'/shots/'+shots[0]['id'],json=body).status_code,409)
        self.assertGreaterEqual(len(self.client.get(url+'/production/history').json()),4)
    def test_timeline_rejects_foreign_media_and_conflicts(self):
        p=self.project();other=self.project();a=self.media(other)
        url='/api/v1/projects/'+p['id']+'/timeline'
        body={'clips':[{'asset_id':a['id'],'duration':1,'start':0}],'version':0}
        self.assertEqual(self.client.put(url,json=body).status_code,422)
        a=self.media(p);body['clips'][0]['asset_id']=a['id']
        self.assertEqual(self.client.put(url,json=body).status_code,200)
        self.assertEqual(self.client.put(url,json=body).status_code,409)
        body['clips'][0]['duration']=-1
        self.assertEqual(self.client.put(url,json=body).status_code,422)
    @unittest.skipUnless(shutil.which('ffmpeg') and shutil.which('ffprobe'),'FFmpeg is required for the real video test')
    def test_real_mp4_export_with_audio(self):
        p=self.project();a=self.media(p);audio=self.media(p,wav(),'audio/wav');url='/api/v1/projects/'+p['id']
        self.assertEqual(self.client.put(url+'/timeline',json={'clips':[{'asset_id':a['id'],'duration':.5},{'asset_id':a['id'],'duration':.5}], 'audio_asset_id':audio['id'],'version':0}).status_code,200)
        response=self.client.post(url+'/renders');self.assertEqual(response.status_code,202)
        row=self.wait_render(p,response.json()['id']);self.assertEqual(row['status'],'completed',row)
        content=self.client.get('/api/v1/assets/'+row['asset_id']+'/content');self.assertEqual(content.status_code,200)
        with main.db() as c:asset=c.execute('SELECT checksum FROM assets WHERE id=?',(row['asset_id'],)).fetchone()
        path=main.DATA/'media'/asset['checksum']
        probe=json.loads(subprocess.check_output(['ffprobe','-v','error','-show_streams','-show_format','-of','json',str(path)]))
        self.assertTrue(any(s['codec_name']=='h264' and s['width']==1280 and s['height']==720 for s in probe['streams']))
        self.assertTrue(any(s['codec_type']=='audio' for s in probe['streams']))
        self.assertAlmostEqual(float(probe['format']['duration']),1,delta=.1)
        self.assertEqual(self.client.post('/api/v1/renders/'+row['id']+'/cancel').status_code,409)
    def test_restore_roundtrip_clones_project_without_overwrite(self):
        p=self.project();a=self.media(p);url='/api/v1/projects/'+p['id']
        self.client.put(url+'/script',json={'text':'Scénario à conserver','version':0})
        self.client.post(url+'/shots',json={'title':'Plan test','asset_id':a['id'],'approved':True})
        self.client.put(url+'/timeline',json={'clips':[{'asset_id':a['id'],'duration':1}],'version':0})
        archive=self.client.get(url+'/export').content
        response=self.client.post('/api/v1/projects/import',files={'file':('project.zip',archive,'application/zip')})
        self.assertEqual(response.status_code,201,response.text)
        restored=response.json();self.assertNotEqual(restored['id'],p['id'])
        state=self.state(restored)
        self.assertEqual(state['script']['text'],'Scénario à conserver')
        self.assertEqual(len(state['shots']),1)
        self.assertNotEqual(state['shots'][0]['asset_id'],a['id'])
        self.assertEqual(state['timeline']['document']['clips'][0]['asset_id'],state['shots'][0]['asset_id'])
        self.assertEqual(self.state(p)['shots'][0]['asset_id'],a['id'])
        self.assertEqual(state['renders'],[])
    def test_restore_rejects_corruption_and_path_traversal(self):
        p=self.project();a=self.media(p)
        original=zipfile.ZipFile(io.BytesIO(self.client.get('/api/v1/projects/'+p['id']+'/export').content))
        for mode in ('corrupt','traversal'):
            out=io.BytesIO()
            with zipfile.ZipFile(out,'w') as archive:
                for name in original.namelist():archive.writestr(name,b'bad' if mode=='corrupt' and name.startswith('media/') else original.read(name))
                if mode=='traversal':archive.writestr('../escape.txt','bad')
            self.assertEqual(self.client.post('/api/v1/projects/import',files={'file':('bad.zip',out.getvalue(),'application/zip')}).status_code,422)
    def test_production_endpoints_are_private(self):
        guest=TestClient(main.app);p=self.project()
        self.assertEqual(guest.get('/api/v1/projects/'+p['id']+'/production').status_code,401)
        self.assertEqual(guest.post('/api/v1/projects/'+p['id']+'/renders').status_code,401)

    def test_docx_extraction_is_local_and_does_not_overwrite_script(self):
        from docx import Document
        document=Document();document.add_paragraph('Texte français.');document.add_paragraph('شخصية أصلية')
        out=io.BytesIO();document.save(out)
        p=self.project();url='/api/v1/projects/'+p['id']
        response=self.client.post(url+'/script/extract',files={'file':('script.docx',out.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        self.assertEqual(response.status_code,200,response.text)
        self.assertIn('شخصية أصلية',response.json()['text'])
        self.assertEqual(self.state(p)['script']['version'],0)
        self.assertEqual(self.client.post(url+'/script/extract',files={'file':('bad.pdf',b'bad','application/pdf')}).status_code,422)

    @unittest.skipUnless(shutil.which('ffmpeg') and shutil.which('ffprobe'),'FFmpeg required')
    def test_failed_render_can_resume_after_missing_media_is_restored(self):
        p=self.project();a=self.media(p);url='/api/v1/projects/'+p['id']
        self.client.put(url+'/timeline',json={'clips':[{'asset_id':a['id'],'duration':.5}],'version':0})
        path=main.DATA/'media'/a['checksum'];content=path.read_bytes();path.unlink()
        rid=self.client.post(url+'/renders').json()['id']
        failed=self.wait_render(p,rid);self.assertEqual(failed['status'],'failed')
        self.assertIn('manquant',failed['error'])
        path.write_bytes(content)
        self.assertEqual(self.client.post('/api/v1/renders/'+rid+'/retry').status_code,200)
        self.assertEqual(self.wait_render(p,rid)['status'],'completed')

    @unittest.skipUnless(shutil.which('ffmpeg') and shutil.which('ffprobe'),'FFmpeg required')
    def test_restart_reuses_verified_segments(self):
        p=self.project();a=self.media(p);url='/api/v1/projects/'+p['id']
        self.client.put(url+'/timeline',json={'clips':[{'asset_id':a['id'],'duration':.5}],'version':0})
        rid=self.client.post(url+'/renders').json()['id']
        self.assertEqual(self.wait_render(p,rid)['status'],'completed')
        segments=list((main.DATA/'renders'/rid).glob('[0-9]*.mp4'))
        self.assertEqual(len(segments),1)
        before=segments[0].stat().st_mtime_ns
        cls=type(self);cls.context.__exit__(None,None,None)
        with main.db() as c:c.execute("UPDATE renders SET status='running',asset_id=NULL,progress=0 WHERE id=?",(rid,))
        cls.context=TestClient(main.app);cls.client=cls.context.__enter__()
        cls.client.post('/api/v1/auth/login',json={'password':'test-only-password-123'})
        self.assertEqual(self.wait_render(p,rid)['status'],'completed')
        self.assertEqual(segments[0].stat().st_mtime_ns,before)

    def test_valid_pdf_extracts_text(self):
        stream=b'BT /F1 16 Tf 30 100 Td (Bonjour Studio) Tj ET'
        objects=[b'<< /Type /Catalog /Pages 2 0 R >>',b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream']
        pdf=bytearray(b'%PDF-1.4\n');offsets=[0]
        for i,obj in enumerate(objects,1):
            offsets.append(len(pdf));pdf.extend(str(i).encode()+b' 0 obj\n'+obj+b'\nendobj\n')
        xref=len(pdf);pdf.extend(b'xref\n0 6\n0000000000 65535 f \n')
        for offset in offsets[1:]:pdf.extend(f'{offset:010d} 00000 n \n'.encode())
        pdf.extend(b'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n'+str(xref).encode()+b'\n%%EOF')
        p=self.project()
        response=self.client.post('/api/v1/projects/'+p['id']+'/script/extract',files={'file':('script.pdf',bytes(pdf),'application/pdf')})
        self.assertEqual(response.status_code,200,response.text)
        self.assertIn('Bonjour Studio',response.json()['text'])
